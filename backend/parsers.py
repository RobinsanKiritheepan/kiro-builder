"""
Kiro Builder — File Parsers
Parse uploaded files: PDF, DOCX, CSV, JSON, ZIP -> structured output.
"""

import io
import os
import json
import zipfile
import tempfile
from pathlib import Path

# ── Supported code extensions for ZIP scanning ────────────────────────
CODE_EXTENSIONS = {
    ".jsx", ".js", ".ts", ".tsx", ".py", ".html", ".css", ".scss",
    ".json", ".md", ".sql", ".yaml", ".yml", ".toml", ".env",
    ".vue", ".svelte", ".php", ".rb", ".go", ".rs", ".java", ".kt",
}


def detect_file_type(filename: str) -> str:
    """Detect parser type from filename extension."""
    ext = Path(filename).suffix.lower()
    mapping = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".csv": "csv",
        ".tsv": "csv",
        ".json": "json",
        ".zip": "zip",
    }
    return mapping.get(ext, "unknown")


def parse_pdf(content: bytes) -> dict:
    """Extract text from a PDF file."""
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        return {"type": "pdf", "content": "", "summary": "PyPDF2 not installed", "suggested_prompt": ""}

    reader = PdfReader(io.BytesIO(content))
    pages_text = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages_text.append(f"--- Page {i + 1} ---\n{text.strip()}")

    full_text = "\n\n".join(pages_text)
    summary = f"PDF with {len(reader.pages)} pages, {len(full_text)} characters extracted."

    return {
        "type": "pdf",
        "content": full_text[:10000],  # cap at 10k chars
        "summary": summary,
        "suggested_prompt": (
            f"Based on this document content, create a web application that presents this information "
            f"in a modern, interactive way:\n\n{full_text[:2000]}"
        ),
    }


def parse_docx(content: bytes) -> dict:
    """Extract text and tables from a DOCX file."""
    try:
        from docx import Document
    except ImportError:
        return {"type": "docx", "content": "", "summary": "python-docx not installed", "suggested_prompt": ""}

    doc = Document(io.BytesIO(content))

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)

    # Extract tables
    tables_text = []
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        tables_text.append(f"Table {i + 1}:\n" + "\n".join(rows))

    full_text = text
    if tables_text:
        full_text += "\n\n" + "\n\n".join(tables_text)

    summary = f"DOCX with {len(paragraphs)} paragraphs and {len(doc.tables)} tables."

    return {
        "type": "docx",
        "content": full_text[:10000],
        "summary": summary,
        "suggested_prompt": (
            f"Based on this document, create a web application that presents this content "
            f"in an engaging, well-structured layout:\n\n{full_text[:2000]}"
        ),
    }


def parse_csv(content: bytes, filename: str = "data.csv") -> dict:
    """Analyze CSV structure and generate a DB schema description."""
    try:
        import pandas as pd
    except ImportError:
        return {"type": "csv", "content": "", "summary": "pandas not installed", "suggested_prompt": ""}

    sep = "\t" if filename.endswith(".tsv") else ","
    df = pd.read_csv(io.BytesIO(content), sep=sep, nrows=100)

    columns_info = []
    for col in df.columns:
        dtype = str(df[col].dtype)
        sample = str(df[col].dropna().iloc[0]) if not df[col].dropna().empty else "N/A"
        nulls = int(df[col].isna().sum())
        columns_info.append(f"  - {col} ({dtype}, nulls: {nulls}, sample: {sample})")

    schema = "\n".join(columns_info)
    preview = df.head(5).to_string(index=False)

    summary = f"CSV with {len(df)} rows and {len(df.columns)} columns: {', '.join(df.columns[:8])}."

    return {
        "type": "csv",
        "content": f"Schema:\n{schema}\n\nPreview (first 5 rows):\n{preview}",
        "summary": summary,
        "suggested_prompt": (
            f"Create a dashboard/CRUD app for this dataset with {len(df.columns)} columns:\n"
            f"{schema}\n\n"
            f"Include: data table with sorting/filtering, charts for numeric columns, "
            f"and a responsive layout."
        ),
    }


def parse_json(content: bytes) -> dict:
    """Analyze JSON structure."""
    try:
        data = json.loads(content)
    except json.JSONDecodeError as e:
        return {"type": "json", "content": "", "summary": f"Invalid JSON: {e}", "suggested_prompt": ""}

    if isinstance(data, list):
        desc = f"JSON array with {len(data)} items"
        if data:
            keys = list(data[0].keys()) if isinstance(data[0], dict) else []
            desc += f", keys: {', '.join(keys[:10])}"
    elif isinstance(data, dict):
        desc = f"JSON object with {len(data)} top-level keys: {', '.join(list(data.keys())[:10])}"
    else:
        desc = f"JSON value: {type(data).__name__}"

    pretty = json.dumps(data, indent=2, ensure_ascii=False)

    return {
        "type": "json",
        "content": pretty[:10000],
        "summary": desc,
        "suggested_prompt": (
            f"Create a web app to display and interact with this JSON data:\n{desc}\n\n"
            f"Preview:\n{pretty[:1500]}"
        ),
    }


def parse_zip(content: bytes) -> dict:
    """Extract and summarize code files from a ZIP archive."""
    files_found = []
    total_size = 0

    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            ext = Path(info.filename).suffix.lower()
            if ext in CODE_EXTENSIONS:
                try:
                    file_content = zf.read(info.filename).decode("utf-8", errors="replace")
                    files_found.append({
                        "name": Path(info.filename).name,
                        "path": info.filename,
                        "content": file_content[:5000],  # cap per file
                        "size": info.file_size,
                    })
                    total_size += info.file_size
                except Exception:
                    pass

    # Sort by path for readability
    files_found.sort(key=lambda f: f["path"])

    file_list = "\n".join(f"  {f['path']} ({f['size']} bytes)" for f in files_found)
    summary = f"ZIP with {len(files_found)} code files, {total_size:,} bytes total."

    # Build architecture overview
    dirs = set()
    for f in files_found:
        parts = Path(f["path"]).parts
        for i in range(1, len(parts)):
            dirs.add("/".join(parts[:i]))

    return {
        "type": "zip",
        "content": f"Files:\n{file_list}\n\nDirectories:\n" + "\n".join(f"  {d}/" for d in sorted(dirs)),
        "summary": summary,
        "files": files_found,
        "suggested_prompt": (
            f"Analyze this project structure and improve/extend it:\n{file_list}"
        ),
    }


def parse_file(filename: str, content: bytes) -> dict:
    """Dispatch to the correct parser based on file extension."""
    file_type = detect_file_type(filename)

    parsers = {
        "pdf":  lambda: parse_pdf(content),
        "docx": lambda: parse_docx(content),
        "csv":  lambda: parse_csv(content, filename),
        "json": lambda: parse_json(content),
        "zip":  lambda: parse_zip(content),
    }

    if file_type in parsers:
        try:
            return parsers[file_type]()
        except Exception as e:
            return {
                "type": file_type,
                "content": "",
                "summary": f"Error parsing {file_type}: {str(e)}",
                "suggested_prompt": "",
            }

    return {
        "type": "unknown",
        "content": content.decode("utf-8", errors="replace")[:5000],
        "summary": f"Unknown file type: {filename}",
        "suggested_prompt": "",
    }
